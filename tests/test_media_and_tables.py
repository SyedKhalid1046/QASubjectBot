"""Unit and integration tests for Images/Media, Calculations, and Tables support."""

import io
import json
import os
import docx
from PIL import Image
import pptx
import pytest

from app import app
from ingestion import (
    DocxLoader,
    DocumentChunk,
    MediaExtractor,
    PptxLoader,
    TokenChunker,
    format_markdown_table,
)
from rag_pipeline import answer_question


@pytest.fixture
def test_client(tmp_path):
    """Creates an isolated test client with dedicated upload, chunks, and faiss directories."""
    app.config["TESTING"] = True
    app.config["UPLOAD_FOLDER"] = str(tmp_path / "documents")
    app.config["CHUNKS_PATH"] = str(tmp_path / "data" / "processed_chunks.json")
    app.config["INDEX_DIR"] = str(tmp_path / "data" / "faiss_index")
    os.makedirs(app.config["UPLOAD_FOLDER"], exist_ok=True)
    os.makedirs(app.config["INDEX_DIR"], exist_ok=True)

    with app.test_client() as client:
        yield client


def test_markdown_table_formatter():
    """Verifies that format_markdown_table generates valid GitHub-Flavored Markdown tables."""
    headers = ["Neuron Type", "Location", "Function"]
    rows = [
        ["Sensory", "PNS", "Afferent signals"],
        ["Motor", "PNS", "Efferent signals"],
        ["Interneuron", "CNS", "Integration"],
    ]
    md = format_markdown_table(headers, rows)
    assert "Neuron Type" in md
    assert "Afferent signals" in md
    assert "Interneuron" in md
    assert "| ---" in md


def test_docx_table_and_image_extraction(tmp_path):
    """Verifies that DocxLoader and MediaExtractor extract tables as Markdown and save embedded pictures."""
    docx_path = str(tmp_path / "neuroscience_lab.docx")
    media_dir = str(tmp_path / "extracted_images")

    # Create dummy image
    img_path = str(tmp_path / "test_synapse.png")
    img = Image.new("RGB", (200, 200), color=(73, 109, 137))
    img.save(img_path)

    doc = docx.Document()
    doc.add_heading("Synaptic Transmission Overview", level=1)
    doc.add_paragraph("Action potentials trigger neurotransmitter release into the synaptic cleft.")
    doc.add_picture(img_path)

    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Neurotransmitter"
    table.cell(0, 1).text = "Primary Effect"
    table.cell(1, 0).text = "GABA"
    table.cell(1, 1).text = "Inhibitory Postsynaptic Potential"

    doc.save(docx_path)

    # 1. Test MediaExtractor
    extracted_images = MediaExtractor.extract_media(docx_path, output_base_dir=media_dir)
    assert 1 in extracted_images
    assert len(extracted_images[1]) >= 1
    assert os.path.exists(os.path.join(media_dir, "neuroscience_lab", extracted_images[1][0]["filename"]))

    # 2. Test DocxLoader loads markdown table
    pages = DocxLoader.load(docx_path)
    assert len(pages) == 1
    page_num, text = pages[0]
    assert page_num == 1
    assert "Neurotransmitter" in text
    assert "Primary Effect" in text
    assert "Inhibitory Postsynaptic Potential" in text


def test_pptx_table_and_image_extraction(tmp_path):
    """Verifies that PptxLoader and MediaExtractor extract tables and slide pictures."""
    pptx_path = str(tmp_path / "brain_imaging.pptx")
    media_dir = str(tmp_path / "extracted_images")

    # Create dummy image
    img_path = str(tmp_path / "mri_scan.png")
    img = Image.new("RGB", (200, 200), color=(120, 50, 200))
    img.save(img_path)

    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Structural Brain Imaging"
    slide.shapes.add_picture(img_path, pptx.util.Inches(1), pptx.util.Inches(1), pptx.util.Inches(3), pptx.util.Inches(3))

    # Add table on slide 2
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    slide2.shapes.title.text = "Modality Comparison"
    table_shape = slide2.shapes.add_table(2, 2, pptx.util.Inches(1), pptx.util.Inches(2), pptx.util.Inches(5), pptx.util.Inches(2))
    table = table_shape.table
    table.cell(0, 0).text = "Modality"
    table.cell(0, 1).text = "Resolution"
    table.cell(1, 0).text = "MRI"
    table.cell(1, 1).text = "High Spatial"

    prs.save(pptx_path)

    # 1. Test MediaExtractor
    extracted_images = MediaExtractor.extract_media(pptx_path, output_base_dir=media_dir)
    assert 1 in extracted_images
    assert len(extracted_images[1]) >= 1
    assert extracted_images[1][0]["page"] == 1

    # 2. Test PptxLoader
    slides = PptxLoader.load(pptx_path)
    assert len(slides) == 2
    assert "Modality" in slides[1][1]
    assert "Resolution" in slides[1][1]
    assert "High Spatial" in slides[1][1]


def test_calculation_and_table_detection_in_chunker():
    """Verifies that TokenChunker accurately detects calculation formulas and tables."""
    chunker = TokenChunker(min_tokens=20, max_tokens=60, target_tokens=40, overlap_tokens=5)
    pages = [
        (
            1,
            "The Nernst equation calculates the equilibrium potential:\n"
            "$$E_{ion} = \\frac{RT}{zF} \\ln \\frac{[ion]_{out}}{[ion]_{in}}$$\n"
            "Where R = 8.314 J/(mol K) and T = 310 K. Calculation yields $E_K = -90 mV$."
        ),
        (
            2,
            "Summary Table:\n\n"
            "| Ion | Extracellular (mM) | Intracellular (mM) |\n"
            "| --- | --- | --- |\n"
            "| Na+ | 145 | 12 |\n"
            "| K+ | 4 | 140 |\n"
        )
    ]
    chunks = chunker.chunk_document_pages(pages, source="membrane_potential.docx")
    assert len(chunks) >= 1

    # Chunk with calculation formulas
    calc_chunks = [c for c in chunks if c.has_calculations]
    assert len(calc_chunks) >= 1

    # Chunk with tables
    table_chunks = [c for c in chunks if c.has_tables]
    assert len(table_chunks) >= 1



def test_rag_pipeline_returns_rich_media_payload():
    """Verifies that answer_question returns image metadata, tables flag, and calculation flags."""
    res = answer_question("What is Magnetic Resonance Imaging (MRI)?", top_k=3, score_threshold=0.05)
    assert res["status"] in ["success", "insufficient_context"]
    assert "images" in res
    assert isinstance(res["images"], list)
    assert "has_images" in res
    assert "has_tables" in res
    assert "has_calculations" in res
    assert "citations" in res


def test_api_documents_metadata_endpoint(test_client):
    """Verifies that /api/documents returns image_count and has_tables per document."""
    res = test_client.get("/api/documents")
    assert res.status_code == 200
    data = res.get_json()
    assert "documents" in data
    for doc in data["documents"]:
        assert "image_count" in doc
        assert "has_tables" in doc
        assert "file_type" in doc
