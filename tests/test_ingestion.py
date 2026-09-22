"""Unit and Integration Tests for QASubjectBot Document Ingestion Pipeline."""

import json
import os
import pytest

from ingestion import (
    DocxLoader,
    DocumentChunk,
    DocumentLoader,
    PDFLoader,
    PptxLoader,
    TextCleaner,
    TextLoader,
    TokenChunker,
    ingest_documents,
    load_processed_chunks,
)
from scripts.generate_corpus import create_pdf


@pytest.fixture
def cleaner():
    return TextCleaner()


@pytest.fixture
def chunker():
    return TokenChunker(
        min_tokens=400,
        max_tokens=1000,
        target_tokens=600,
        overlap_tokens=100,
    )


@pytest.fixture
def docs_dir():
    return os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "documents"))


def test_cleaner_removes_page_numbers_and_footers(cleaner):
    raw_text = (
        "QASubjectBot Technical Knowledge Series\n"
        "Section 1.1: Foundations of Deep Learning\n"
        "Modern deep neural networks require careful trans-\nformer initialization.\n"
        "Page 1 of 5\n"
        "Confidential & Proprietary - AI Research Lab Reference Corpus\n"
    )
    cleaned = cleaner.clean(raw_text)

    # Assertions
    assert "QASubjectBot Technical Knowledge Series" not in cleaned
    assert "Page 1 of 5" not in cleaned
    assert "Confidential & Proprietary" not in cleaned
    assert "transformer initialization" in cleaned  # Tests hyphen repair
    assert "Section 1.1: Foundations of Deep Learning" in cleaned


def test_cleaner_handles_empty_and_whitespace(cleaner):
    assert cleaner.clean("") == ""
    assert cleaner.clean("   \n\n\t\r\n   ") == ""


def test_pdf_loader_isolated(tmp_path):
    sample_pdf = os.path.join(tmp_path, "sample_test.pdf")
    doc_info = {
        "filename": "sample_test.pdf",
        "title": "Sample PDF Document",
        "pages": [
            [
                ("Section 1: Introduction", "This is page one of our unit test PDF document for PDFLoader verification."),
            ],
            [
                ("Section 2: Calculations", "This is page two containing second page text information for validation."),
            ],
        ],
    }
    create_pdf(doc_info, sample_pdf)
    assert os.path.exists(sample_pdf)

    pages = PDFLoader.load(sample_pdf)
    assert len(pages) == 2
    page_num, text = pages[0]
    assert page_num == 1
    assert "unit test PDF" in text


def test_text_loader_isolated(tmp_path):
    sample_txt = os.path.join(tmp_path, "sample_test.txt")
    with open(sample_txt, "w", encoding="utf-8") as f:
        f.write("ACCOUNTING PRINCIPLES\n\nThis is a sample plain text file.")

    pages = TextLoader.load(sample_txt)
    assert len(pages) == 1
    assert pages[0][0] == 1
    assert "ACCOUNTING PRINCIPLES" in pages[0][1]


def test_docx_loader(tmp_path):
    import docx
    sample_docx = os.path.join(tmp_path, "sample_test.docx")
    doc = docx.Document()
    doc.add_heading("Biopsychology Overview", level=1)
    doc.add_paragraph("Biopsychology is the scientific study of the biology of behavior.")
    doc.save(sample_docx)
    assert os.path.exists(sample_docx)

    pages = DocxLoader.load(sample_docx)
    assert len(pages) == 1
    assert pages[0][0] == 1
    assert "BIOPSYCHOLOGY" in pages[0][1].upper()


def test_pptx_loader(docs_dir, tmp_path):
    pptx_files = [f for f in os.listdir(docs_dir) if f.endswith(".pptx")]
    if pptx_files:
        sample_pptx = os.path.join(docs_dir, pptx_files[0])
    else:
        import pptx
        sample_pptx = os.path.join(tmp_path, "sample_test.pptx")
        prs = pptx.Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[0])
        s.shapes.title.text = "Biopsychology Test Presentation"
        prs.save(sample_pptx)

    assert os.path.exists(sample_pptx)

    slides = PptxLoader.load(sample_pptx)
    assert len(slides) >= 1
    assert slides[0][0] == 1
    assert len(slides[0][1]) > 0


def test_chunker_bounds_and_token_counts(chunker):
    text = (
        "Marginal costing is a technique of cost accounting that focuses on variable costs and contribution margin. "
        * 40
    )
    pages_cleaned = [(1, text)]
    chunks = chunker.chunk_document_pages(pages_cleaned, source="test.docx")

    assert len(chunks) >= 1
    for chunk in chunks:
        assert isinstance(chunk, DocumentChunk)
        assert chunk.token_count <= 1000
        assert chunk.token_count >= 100
        assert chunk.source == "test.docx"
        assert chunk.page == 1
        assert "test.docx:p1:c" in chunk.chunk_id


def test_chunk_metadata_schema_and_serialization(tmp_path, docs_dir):
    output_json = os.path.join(tmp_path, "test_chunks.json")
    chunks = ingest_documents(
        docs_dir=docs_dir,
        output_path=output_json,
        min_tokens=400,
        max_tokens=1000,
        target_tokens=600,
        overlap_tokens=100,
    )

    # 1. Verify chunk count from current documents
    assert len(chunks) >= 30
    assert os.path.exists(output_json)

    # 2. Verify all chunks adhere to metadata invariants
    for idx, c in enumerate(chunks):
        assert len(c.text.strip()) > 0
        assert c.source.endswith((".pdf", ".txt", ".docx", ".pptx"))
        assert isinstance(c.page, int) and c.page >= 1
        assert isinstance(c.chunk_index, int) and c.chunk_index >= 0
        assert isinstance(c.token_count, int) and c.token_count > 0
        assert isinstance(c.metadata.get("char_count"), int)
        assert c.chunk_id == f"{c.source}:p{c.page}:c{c.chunk_index}"

    # 3. Verify deserialization
    loaded = load_processed_chunks(output_json)
    assert len(loaded) == len(chunks)
    assert loaded[0].chunk_id == chunks[0].chunk_id
    assert loaded[0].text == chunks[0].text

